"""
상담 세션 저장소 (Session Repository)

한 상담 = 한 줄. 메인 시트에는 사람이 읽고 검색·분류할 칸(날짜·일시·PB·고객·상태
+ 7요인 요약)을 펼치고, 전체 AnalysisResult는 마지막 칸에 JSON 문자열로 저장한다.
이렇게 하면 PB가 엑셀을 열어도 한눈에 보이고, 프로그램은 JSON 칸으로 완전 복원한다.

고객 저장소와 동일하게 추상 인터페이스로 분리 → 나중에 Supabase 교체 가능.
모든 상담은 customer_id로 묶인다(이름 아님 → 동명이인 안전).
"""

from __future__ import annotations

import json
import os
import random
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from datetime import datetime
from typing import Optional

from openpyxl import Workbook, load_workbook

from models import SessionStatus, FactorStatus, Confidence
from analysis_schema import AnalysisResult


# ---------------------------------------------------------------------------
# 세션 데이터 구조
# ---------------------------------------------------------------------------

@dataclass
class Session:
    session_id: str
    customer_id: str          # 어떤 고객인지 (진짜 식별자)
    customer_name: str        # 표시용
    pb_name: str
    consult_date: str         # "YYYY-MM-DD"
    consult_datetime: str     # ISO 일시
    status: str               # SessionStatus 값
    raw_text: str             # 고객 상담 원문
    result: AnalysisResult    # 전체 분석 결과

    # --- 메인 시트에 펼칠 7요인 요약 생성 ---
    def summary(self) -> dict:
        r = self.result
        gr = ""
        if r.goal_return.return_min is not None:
            lo = r.goal_return.return_min * 100
            hi = (r.goal_return.return_max or r.goal_return.return_min) * 100
            gr = f"{lo:.0f}~{hi:.0f}%" if hi != lo else f"{lo:.0f}%"
        risk = ""
        b = r.risk_tolerance.binding
        if b == "willingness":
            risk = r.risk_tolerance.willingness.level or ""
        elif b == "capacity":
            risk = r.risk_tolerance.capacity.level or ""
        unresolved = sum(1 for f in r.flags if not f.resolved)
        return {
            "목표수익률": gr,
            "위험등급": risk,
            "투자기간(년)": r.horizon.years if r.horizon.years is not None else "",
            "미해결플래그수": unresolved,
            "추가질문수": len(r.follow_up_questions),
        }


# ---------------------------------------------------------------------------
# 추상 인터페이스
# ---------------------------------------------------------------------------

class SessionRepository(ABC):
    @abstractmethod
    def save(self, customer_id: str, customer_name: str, pb_name: str,
             raw_text: str, result: AnalysisResult,
             consult_date: Optional[str] = None,
             status: str = SessionStatus.DRAFT.value) -> Session: ...

    @abstractmethod
    def get(self, session_id: str) -> Optional[Session]: ...

    @abstractmethod
    def update_status(self, session_id: str, status: str) -> None: ...

    @abstractmethod
    def list_by_customer(self, customer_id: str) -> list[Session]: ...

    @abstractmethod
    def list_filtered(self, pb_name: Optional[str] = None,
                      date_from: Optional[str] = None,
                      date_to: Optional[str] = None,
                      status: Optional[str] = None) -> list[Session]: ...


# ---------------------------------------------------------------------------
# 엑셀 구현
# ---------------------------------------------------------------------------

class ExcelSessionRepository(SessionRepository):
    SHEET = "Sessions"
    # 펼친 칸 + 마지막에 result_json
    HEADERS = ["session_id", "customer_id", "customer_name", "pb_name",
               "consult_date", "consult_datetime", "status",
               "목표수익률", "위험등급", "투자기간(년)", "미해결플래그수", "추가질문수",
               "raw_text", "result_json"]

    def __init__(self, path: str):
        self.path = path
        if not os.path.exists(path):
            wb = Workbook(); ws = wb.active; ws.title = self.SHEET
            ws.append(self.HEADERS); wb.save(path)

    def _load(self):
        wb = load_workbook(self.path)
        if self.SHEET not in wb.sheetnames:
            ws = wb.create_sheet(self.SHEET); ws.append(self.HEADERS)
        return wb, wb[self.SHEET]

    def _rows(self) -> list[dict]:
        wb, ws = self._load()
        out, header = [], None
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                header = list(row); continue
            if row is None or all(c is None for c in row):
                continue
            out.append(dict(zip(header, row)))
        return out

    def _gen_id(self, existing: set[str]) -> str:
        n = 0
        for s in existing:
            if s and s.startswith("S") and s[1:].isdigit():
                n = max(n, int(s[1:]))
        cand = f"S{n+1:05d}"
        while cand in existing:
            cand = f"S{random.randint(1,99999):05d}"
        return cand

    def _row_to_session(self, d: dict) -> Session:
        result = AnalysisResult.from_dict(json.loads(d["result_json"]))
        return Session(
            session_id=str(d["session_id"]), customer_id=str(d["customer_id"]),
            customer_name=str(d["customer_name"]), pb_name=str(d["pb_name"]),
            consult_date=str(d["consult_date"]), consult_datetime=str(d["consult_datetime"]),
            status=str(d["status"]), raw_text=str(d.get("raw_text") or ""), result=result,
        )

    def next_id(self) -> str:
        """다음 session_id를 미리 확보한다 (백그라운드 분석용)."""
        return self._gen_id({str(r["session_id"]) for r in self._rows()})

    def update_result(self, session_id: str, result: AnalysisResult) -> None:
        """분석 결과(JSON + 요약 칸)를 덮어쓴다. 플래그 해소 등에 사용."""
        wb, ws = self._load()
        header = [c.value for c in ws[1]]
        col = {h: i for i, h in enumerate(header)}

        gr = ""
        if result.goal_return.return_min is not None:
            lo = result.goal_return.return_min * 100
            hi = (result.goal_return.return_max or result.goal_return.return_min) * 100
            gr = f"{lo:.0f}~{hi:.0f}%" if hi != lo else f"{lo:.0f}%"
        risk = ""
        b = result.risk_tolerance.binding
        if b == "willingness":
            risk = result.risk_tolerance.willingness.level or ""
        elif b == "capacity":
            risk = result.risk_tolerance.capacity.level or ""

        for row in ws.iter_rows(min_row=2):
            if row[col["session_id"]].value == session_id:
                row[col["result_json"]].value = json.dumps(result.to_dict(), ensure_ascii=False)
                row[col["목표수익률"]].value = gr
                row[col["위험등급"]].value = risk
                row[col["투자기간(년)"]].value = result.horizon.years if result.horizon.years is not None else ""
                row[col["미해결플래그수"]].value = sum(1 for f in result.flags if not f.resolved)
                row[col["추가질문수"]].value = len(result.follow_up_questions)
                wb.save(self.path)
                return
        raise KeyError(f"세션을 찾을 수 없습니다: {session_id}")

    def save(self, customer_id, customer_name, pb_name, raw_text, result,
             consult_date=None, status=SessionStatus.DRAFT.value,
             session_id: Optional[str] = None) -> Session:
        now = datetime.now()
        rows = self._rows()
        sid = session_id if session_id else self._gen_id({str(r["session_id"]) for r in rows})
        sess = Session(
            session_id=sid, customer_id=customer_id, customer_name=customer_name,
            pb_name=pb_name, consult_date=consult_date or now.strftime("%Y-%m-%d"),
            consult_datetime=now.isoformat(timespec="seconds"), status=status,
            raw_text=raw_text, result=result,
        )
        s = sess.summary()
        wb, ws = self._load()
        ws.append([sess.session_id, sess.customer_id, sess.customer_name, sess.pb_name,
                   sess.consult_date, sess.consult_datetime, sess.status,
                   s["목표수익률"], s["위험등급"], s["투자기간(년)"], s["미해결플래그수"], s["추가질문수"],
                   sess.raw_text, json.dumps(result.to_dict(), ensure_ascii=False)])
        wb.save(self.path)
        return sess

    def get(self, session_id: str) -> Optional[Session]:
        for d in self._rows():
            if str(d["session_id"]) == session_id:
                return self._row_to_session(d)
        return None

    def update_status(self, session_id: str, status: str) -> None:
        wb, ws = self._load()
        header = [c.value for c in ws[1]]
        sid_col = header.index("session_id") + 1
        st_col = header.index("status") + 1
        for row in ws.iter_rows(min_row=2):
            if row[sid_col-1].value == session_id:
                row[st_col-1].value = status
                wb.save(self.path); return
        raise KeyError(f"세션을 찾을 수 없습니다: {session_id}")

    def list_by_customer(self, customer_id: str) -> list[Session]:
        return [self._row_to_session(d) for d in self._rows()
                if str(d["customer_id"]) == customer_id]

    def list_filtered(self, pb_name=None, date_from=None, date_to=None, status=None) -> list[Session]:
        out = []
        for d in self._rows():
            if pb_name and str(d["pb_name"]) != pb_name: continue
            cd = str(d["consult_date"])
            if date_from and cd < date_from: continue
            if date_to and cd > date_to: continue
            if status and str(d["status"]) != status: continue
            out.append(self._row_to_session(d))
        return out
